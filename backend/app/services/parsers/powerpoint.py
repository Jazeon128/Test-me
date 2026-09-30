from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from .base_parser import BaseParser, ParsedDocument, ParsedSection


class PowerPointParser(BaseParser):
    def _shape_text(self, shapes):
        """Read text recursively, including grouped shapes and tables."""
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self._shape_text(shape.shapes)
            elif shape.has_table is True:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            yield cell.text.strip()
            elif hasattr(shape, "text") and shape.text.strip():
                yield shape.text.strip()

    def parse(self, file_path: str) -> ParsedDocument:
        """
        Parses a PowerPoint file and returns the text from all slides.
        Includes text from shapes and notes slides.
        """
        try:
            prs = Presentation(file_path)
            sections = []
            full_text_parts = []

            for i, slide in enumerate(prs.slides):
                # Extract text from shapes in the slide
                slide_text = list(self._shape_text(slide.shapes))

                # Extract text from notes slide (if available)
                if slide.has_notes_slide:
                    notes_slide = slide.notes_slide
                    slide_text.extend(self._shape_text(notes_slide.shapes))

                if slide_text:
                    slide_content = "\n".join(slide_text)
                    full_text_parts.append(slide_content)
                    sections.append(
                        ParsedSection(
                            text=slide_content,
                            page=i + 1,  # Use slide number as page number
                            section=f"Slide {i + 1}",
                        )
                    )

            full_text = "\n\n".join(full_text_parts)

            return ParsedDocument(
                full_text=full_text,
                sections=sections,
                num_pages=len(prs.slides),
                metadata={"source": file_path, "type": "pptx"},
            )
        except Exception as e:
            raise Exception(f"Failed to parse PowerPoint file: {str(e)}")
