from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.services.parsers.docx_parser import DOCXParser
from app.services.parsers.html_parser import HTMLParser
from app.services.parsers.powerpoint import PowerPointParser


def test_docx_body_order_and_short_table_rows(tmp_path):
    doc = Document()
    doc.add_heading("Document heading", level=1)
    before = "The paragraph before the table has useful content."
    after = "The paragraph after the table has more useful content."
    doc.add_paragraph(before)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    doc.add_paragraph(after)
    path = tmp_path / "ordered.docx"
    doc.save(path)

    result = DOCXParser().parse(str(path))

    assert [section.text for section in result.sections] == [before, "A | B", after]
    assert result.full_text == f"{before} A | B {after}"
    assert result.title == "Document heading"
    assert all(section.section == result.title for section in result.sections)
    for section in result.sections:
        assert result.full_text[section.start_char:section.end_char] == section.text


def test_powerpoint_groups_nested_groups_and_tables(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Grouped text"
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Nested text"
    table = slide.shapes.add_table(1, 2, 0, 0, Inches(2), Inches(1)).table
    table.cell(0, 0).text = "First cell"
    table.cell(0, 1).text = "Second cell"
    path = tmp_path / "groups.pptx"
    prs.save(path)

    result = PowerPointParser().parse(str(path))

    assert result.full_text.splitlines() == [
        "Grouped text", "Nested text", "First cell", "Second cell",
    ]
    assert result.sections[0].page == 1
    assert result.num_pages == 1


def test_html_declared_windows_1252(tmp_path):
    path = tmp_path / "encoded.html"
    text = "The café serves coffee throughout the morning."
    path.write_bytes(
        f'<html><head><meta charset="windows-1252"></head>'
        f'<body><p>{text}</p></body></html>'.encode("windows-1252")
    )

    assert HTMLParser().parse(str(path)).full_text == text


def test_html_nested_blocks_and_headings(tmp_path):
    path = tmp_path / "nested.html"
    first = "This passage appears exactly once in the document."
    second = "Another passage belongs to the second heading."
    path.write_bytes(
        f"<div><h1>First</h1><div><p>{first}</p></div>"
        f"<h2>Second</h2><div><p>{second}</p></div></div>".encode()
    )

    result = HTMLParser().parse(str(path))

    assert [s.text for s in result.sections] == [first, second]
    assert [s.section for s in result.sections] == ["First", "Second"]
    assert result.full_text == f"{first} {second}"


def test_html_own_text_nested_lists_and_cells(tmp_path):
    path = tmp_path / "mixed.html"
    texts = [
        "Direct container text must also be preserved.",
        "Parent list item with inline emphasis included.",
        "Nested list item must appear only once.",
        "First table cell contains a nested paragraph.",
        "Second table cell contains its own text.",
    ]
    path.write_bytes(
        f"<div>{texts[0]}<ul><li>Parent list item with <em>inline emphasis</em>"
        f" included.<ul><li>{texts[2]}</li></ul></li></ul>"
        f"<table><tr><td><p>{texts[3]}</p></td><td>{texts[4]}</td></tr>"
        "</table></div>".encode()
    )

    result = HTMLParser().parse(str(path))

    assert [s.text for s in result.sections] == texts
    assert result.full_text == " ".join(texts)


def test_html_realistic_page_keeps_visible_content():
    path = Path(__file__).parents[1] / "fixtures" / "html_check.html"
    texts = [
        "Amazon S3 stores objects in buckets and is designed for durability.",
        "Quoted guidance about lifecycle policies moving objects to Glacier after thirty days.",
        "aws s3 cp local.txt s3://bucket/remote.txt --storage-class STANDARD_IA",
        "Loose body text that sits directly inside the article element without any wrapper tag.",
        "A normal paragraph with bold words and a link inside it that should stay whole.",
    ]

    result = HTMLParser().parse(str(path))

    assert [s.text for s in result.sections] == texts
    assert all(s.section == "Guide" for s in result.sections)
    assert result.full_text == " ".join(texts)
    assert "Home" not in result.full_text
    assert "Docs" not in result.full_text


@pytest.mark.parametrize("owner", [
    "section", "article", "main", "blockquote", "pre", "dd", "dt", "figcaption", "body",
])
def test_html_additional_owners_preserve_text_once(tmp_path, owner):
    path = tmp_path / "owners.html"
    before = "Visible text before the nested paragraph must stay."
    nested = "Visible nested paragraph must appear exactly once."
    after = "Visible text after the nested paragraph must stay."
    content = f"<{owner}>{before}<p>{nested}</p>{after}</{owner}>"
    if owner != "body":
        content = f"<body>{content}</body>"
    path.write_bytes(content.encode())

    result = HTMLParser().parse(str(path))

    assert [s.text for s in result.sections] == [before, nested, after]
    assert result.full_text == f"{before} {nested} {after}"


def test_html_excluded_elements_and_short_blocks(tmp_path):
    path = tmp_path / "excluded.html"
    excluded = "Excluded text is long enough to pass the paragraph filter."
    kept = "A visible paragraph must still be preserved in this document."
    elements = "".join(
        f"<{tag}>{excluded}</{tag}>"
        for tag in (
            "script", "style", "noscript", "template", "nav", "header", "footer", "aside",
        )
    )
    path.write_bytes(
        f"<body><h1>Guide</h1>{elements}<p>12345678901234567890</p>"
        f"<p>Short fragment</p><p>{kept}</p></body>".encode()
    )

    result = HTMLParser().parse(str(path))

    assert result.full_text == kept
    assert [s.text for s in result.sections] == [kept]
    assert result.sections[0].section == "Guide"
